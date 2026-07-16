"""Tiny read-only pandas compatibility layer for legacy XLSX validation.

It intentionally implements only the operations used by
``generate_welds.load_frozen_weld_instance``.  Formal experiment code never
imports this module.
"""

from math import isnan

from openpyxl import load_workbook


class DataFrame:
    def __init__(self, rows):
        self._rows = list(rows)
        self.columns = list(self._rows[0]) if self._rows else []

    def iterrows(self):
        return enumerate(self._rows)

    def sort_values(self, column, kind=None):
        del kind
        return DataFrame(sorted(self._rows, key=lambda row: row[column]))


def isna(value):
    return value is None or (isinstance(value, float) and isnan(value))


def read_excel(path, sheet_name=None, engine=None):
    del engine
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    for worksheet in workbook.worksheets:
        values = worksheet.iter_rows(values_only=True)
        headers = list(next(values))
        rows = [dict(zip(headers, row)) for row in values]
        result[worksheet.title] = DataFrame(rows)
    if sheet_name is None:
        return result
    return result[sheet_name]
